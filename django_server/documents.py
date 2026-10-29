import math
import re
import threading
import unicodedata
from collections import Counter
from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile

from docx import Document
from pypdf import PdfReader
from pypdf.errors import PdfReadError
from docx.opc.exceptions import PackageNotFoundError

from django_server.chat_errors import ChatError


BASE_DIR = Path(__file__).resolve().parent.parent
KNOWLEDGE_DIR = BASE_DIR / 'knowledge_base'
SUPPORTED_EXTENSIONS = {'.pdf', '.docx', '.txt', '.md', '.markdown'}
MAX_DOCUMENT_BYTES = 10 * 1024 * 1024
MAX_TOTAL_DOCUMENT_BYTES = 25 * 1024 * 1024
MAX_DOCUMENT_COUNT = 50
CHUNK_SIZE = 1200
CHUNK_OVERLAP = 160
_INDEX_LOCK = threading.Lock()
_INDEX_SIGNATURE = None
_INDEX_CHUNKS = []


def _read_document(path):
    try:
        content = path.read_bytes()
        if path.suffix.lower() == '.pdf':
            reader = PdfReader(BytesIO(content), strict=False)
            if reader.is_encrypted:
                raise ChatError(f'Tài liệu {path.name} được mã hóa, không thể trích xuất.', 400)
            pages = []
            for page_number, page in enumerate(reader.pages, start=1):
                text = page.extract_text() or ''
                if text.strip():
                    pages.append(f'[Trang {page_number}]\n{text}')
            return '\n\n'.join(pages)
        if path.suffix.lower() == '.docx':
            document = Document(BytesIO(content))
            paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
            for table in document.tables:
                paragraphs.extend(
                    ' | '.join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    for row in table.rows
                )
            return '\n'.join(paragraphs)
        return content.decode('utf-8-sig')
    except ChatError:
        raise
    except (
        OSError,
        UnicodeError,
        ValueError,
        BadZipFile,
        PdfReadError,
        PackageNotFoundError,
    ) as error:
        raise ChatError(
            f'Không đọc được tài liệu {path.name}. Hãy kiểm tra tệp có bị hỏng hoặc đúng định dạng.',
            400,
        ) from error


def _split_chunks(text, source):
    text = re.sub(r'\s+', ' ', text).strip()
    chunks = []
    start = 0
    while start < len(text):
        end = min(start + CHUNK_SIZE, len(text))
        if end < len(text):
            boundary = text.rfind(' ', start + CHUNK_SIZE // 2, end)
            if boundary > start:
                end = boundary
        chunk = text[start:end].strip()
        if chunk:
            chunks.append({'source': source, 'text': chunk})
        if end >= len(text):
            break
        start = max(start + 1, end - CHUNK_OVERLAP)
    return chunks


def _tokens(text):
    normalized = unicodedata.normalize('NFD', text.lower())
    normalized = ''.join(char for char in normalized if unicodedata.category(char) != 'Mn')
    normalized = normalized.replace('đ', 'd')
    return re.findall(r'[a-z0-9]+', normalized)


def _load_index():
    global _INDEX_SIGNATURE, _INDEX_CHUNKS
    KNOWLEDGE_DIR.mkdir(parents=True, exist_ok=True)
    files = sorted(
        path for path in KNOWLEDGE_DIR.iterdir()
        if path.suffix.lower() in SUPPORTED_EXTENSIONS and not path.is_symlink()
    )
    if len(files) > MAX_DOCUMENT_COUNT:
        raise ChatError(f'Thư mục tài liệu vượt quá {MAX_DOCUMENT_COUNT} tệp.', 413)

    signature = []
    total_size = 0
    for path in files:
        try:
            stat = path.stat()
        except OSError as error:
            raise ChatError(f'Không thể kiểm tra tài liệu {path.name}.', 400) from error
        if stat.st_size > MAX_DOCUMENT_BYTES:
            raise ChatError(f'Tài liệu {path.name} vượt quá giới hạn 10 MB.', 413)
        total_size += stat.st_size
        signature.append((path.name, stat.st_mtime_ns, stat.st_size))
    if total_size > MAX_TOTAL_DOCUMENT_BYTES:
        raise ChatError('Tổng dung lượng tài liệu vượt quá giới hạn 25 MB.', 413)

    signature = tuple(signature)
    with _INDEX_LOCK:
        if signature == _INDEX_SIGNATURE:
            return _INDEX_CHUNKS
        chunks = []
        for path in files:
            chunks.extend(_split_chunks(_read_document(path), path.name))
        _INDEX_SIGNATURE = signature
        _INDEX_CHUNKS = chunks
        return chunks


def find_relevant_passages(query, limit=3):
    chunks = _load_index()
    query_tokens = set(_tokens(query))
    if not chunks or not query_tokens:
        return []

    tokenized_chunks = [_tokens(chunk['text']) for chunk in chunks]
    document_frequency = Counter(
        token for tokens in tokenized_chunks for token in set(tokens)
    )
    average_length = max(1, sum(map(len, tokenized_chunks)) / len(tokenized_chunks))
    scored = []
    for chunk, tokens in zip(chunks, tokenized_chunks):
        frequencies = Counter(tokens)
        score = 0.0
        for token in query_tokens.intersection(frequencies):
            inverse_frequency = math.log(
                1 + (len(chunks) - document_frequency[token] + 0.5)
                / (document_frequency[token] + 0.5)
            )
            frequency = frequencies[token]
            length_normalization = 1.2 * (0.25 + 0.75 * len(tokens) / average_length)
            score += inverse_frequency * (frequency * 2.2) / (frequency + length_normalization)
        if score > 0:
            scored.append((score, chunk))

    scored.sort(key=lambda result: result[0], reverse=True)
    return [chunk for _, chunk in scored[:limit]]
